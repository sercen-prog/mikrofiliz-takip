import json
import datetime
import io
import openpyxl
import streamlit as st
import gspread
import pandas as pd
from google.oauth2.service_account import Credentials

# Reportlab PDF Kütüphaneleri
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

st.set_page_config(
    page_title="🌱 Mikrofiliz Üretim, Satış, Fire ve Tohum Takip",
    page_icon="🌱",
    layout="wide"
)

# 🔒 ŞİFRE KORUMASI VE OTURUM KONTROLÜ
APP_PASSWORD = "tarlam2026"

def check_password():
    if "authenticated" not in st.session_state:
        st.session_state.authenticated = False

    if not st.session_state.authenticated:
        st.markdown("<br><br>", unsafe_allow_html=True)
        col1, col2, col3 = st.columns([1, 1.5, 1])
        with col2:
            st.markdown("### 🔒 Güvenli Giriş Paneli")
            pwd_input = st.text_input("Sisteme erişmek için şifrenizi girin:", type="password", key="login_password_input")
            if st.button("🔓 Giriş Yap", use_container_width=True):
                if pwd_input == APP_PASSWORD:
                    st.session_state.authenticated = True
                    st.success("Giriş başarılı! Yönlendiriliyorsunuz...")
                    st.rerun()
                else:
                    st.error("Hatalı şifre! Lütfen tekrar deneyin.")
        return False
    return True

if not check_password():
    st.stop()

# Orijinal Excel Şablon Verileri
ORIGINAL_PRODUCT_TEMPLATES = [
    ("Amaranth", "10", "3", 5),
    ("Bezelye", "12", "4", 150),
    ("Fesleğen İtalyan (İntfa)", "18", "5", 10),
    ("Havuç", "22", "5", 8),
    ("Kişniş", "22", "5", 25),
    ("Reyhan", "18", "5", 10),
    ("Mor Lahana", "9", "3", 12),
    ("Pembe Turp", "9", "3", 30),
    ("Rambo Turp Mor", "9", "3", 30),
    ("Roka", "9", "3", 12),
    ("Brokoli", "9", "3", 12),
    ("Hardal (Mizuna)", "9", "3", 12),
    ("Beyaz Turp", "9", "3", 30),
    ("Fesleğen (Thai)", "18", "5", 10),
    ("Hodan", "12", "3", 40),
    ("Çemen", "10", "3", 35),
    ("Frenk soğanı", "14", "5", 15),
    ("Kırmızı Pancar", "12", "3", 40),
    ("Maş Fasulyesi", "10", "3", 100),
    ("Marul", "12", "3", 8),
    ("Kuzu kulağı", "18", "5", 8),
    ("Kırmızı Damarlı Kuzukulağı", "18", "5", 8),
    ("Ayçekirdeği", "12", "3", 120),
    ("Mor Alabaş", "10", "3", 12),
    ("Kara lahana", "10", "3", 12),
    ("Tere", "8", "3", 12)
]

DEFAULT_CUSTOMERS = [
    "Le Meridyen", "Damlıca Çiftliği", "The Crepe Escape", "Manca", "Mandarin"
]

HEADERS = [
    "Ürün Adı", "Üretim Adeti", "Stoktan Çıkan (Satılan)", "Çöpe Atılan (Fire)", 
    "Kalan Ürün", "Satış Yapılan Yer / Müşteri", "Ekim Günü", "Hasat Süresi (gün)", 
    "Hasat Tarihi", "Karanlıkta (gün)", "Işığa Çıkış Tarihi", "Not"
]

def safe_int(val):
    if not val: return 0
    try:
        val_str = str(val).replace("💡", "").strip()
        return int(float(val_str))
    except:
        return 0

def safe_float(val):
    if not val: return 0.0
    try:
        val_str = str(val).replace("💡", "").strip()
        return float(val_str)
    except:
        return 0.0

def parse_tr_date(dt_str):
    if not dt_str: return None
    dt_str = str(dt_str).replace("💡", "").strip()
    try:
        if "." in dt_str:
            p = dt_str.split(".")
            return datetime.date(int(p[2]), int(p[1]), int(p[0]))
        elif "-" in dt_str:
            p = dt_str.split("-")
            if len(p[0]) == 4: return datetime.date(int(p[0]), int(p[1]), int(p[2]))
            else: return datetime.date(int(p[2]), int(p[1]), int(p[0]))
    except: return None
    return None

def format_date_tr(dt_val):
    if not dt_val: return ""
    if isinstance(dt_val, (datetime.date, datetime.datetime)):
        return dt_val.strftime("%d.%m.%Y")
    parsed = parse_tr_date(dt_val)
    if parsed:
        return parsed.strftime("%d.%m.%Y")
    return str(dt_val)

def fix_tr_chars_for_pdf(txt):
    if not txt: return ""
    txt = str(txt)
    tr_map = {
        'ğ': 'g', 'Ğ': 'G', 'ı': 'i', 'İ': 'I',
        'ö': 'o', 'Ö': 'O', 'ü': 'u', 'Ü': 'U',
        'ş': 's', 'Ş': 'S', 'ç': 'c', 'Ç': 'C'
    }
    for k, v in tr_map.items():
        txt = txt.replace(k, v)
    return txt

def get_clean_empty_sheet_data():
    rows = [HEADERS]
    for name, hasat, kar, _ in ORIGINAL_PRODUCT_TEMPLATES:
        rows.append([name, "0", "0", "0", "0", "", "", hasat, "", kar, "", ""])
    return {"Parti No.1": rows}

# Google Sheets Bağlantı Kurulumu
@st.cache_resource
def get_gspread_client():
    scopes = [
        "https://www.googleapis.com/auth/spreadsheets",
        "https://www.googleapis.com/auth/drive"
    ]
    if "gcp_json" in st.secrets:
        creds_dict = json.loads(st.secrets["gcp_json"])
        creds = Credentials.from_service_account_info(creds_dict, scopes=scopes)
    elif "gcp_service_account" in st.secrets:
        creds_dict = dict(st.secrets["gcp_service_account"])
        creds = Credentials.from_service_account_info(creds_dict, scopes=scopes)
    else:
        raise ValueError("Google Service Account anahtarı Streamlit Secrets içinde bulunamadı!")
    return gspread.authorize(creds)

def load_data_from_sheets():
    try:
        gc = get_gspread_client()
        sh = gc.open("Mikrofiliz_Verileri")
        sheets = {}
        template_dict = {p[0]: (p[1], p[2]) for p in ORIGINAL_PRODUCT_TEMPLATES}

        for ws in sh.worksheets():
            if ws.title == "Tohum_Stoklari":
                continue
            rows = ws.get_all_values()
            if len(rows) > 1:
                clean_prod = []
                for r in rows[1:]:
                    if len(r) > 0 and str(r[0]).strip() != "Toplam Üretim":
                        u_name = str(r[0]).strip()
                        u_adet = safe_int(r[1]) if len(r)>1 else 0
                        s_cikan = safe_int(r[2]) if len(r)>2 else 0
                        fire = safe_int(r[3]) if len(r)>3 else 0
                        kalan = safe_int(r[4]) if len(r)>4 else 0
                        musteri = str(r[5]).strip() if len(r)>5 and not str(r[5]).strip().isdigit() else ""
                        ekim_str = format_date_tr(r[6]) if len(r)>6 else ""
                        
                        def_hasat, def_kar = template_dict.get(u_name, ("10", "3"))
                        has_suresi = str(r[7]).strip() if len(r)>7 and str(r[7]).strip().isdigit() else def_hasat
                        kar_suresi = str(r[9]).strip() if len(r)>9 and str(r[9]).strip().isdigit() else def_kar

                        has_tarihi = format_date_tr(r[8]) if len(r)>8 else ""
                        isik_tarihi = format_date_tr(r[10]) if len(r)>10 else ""
                        not_txt = str(r[11]).strip() if len(r)>11 else ""

                        clean_prod.append([
                            u_name, str(u_adet), str(s_cikan), str(fire), str(kalan),
                            musteri, ekim_str, has_suresi, has_tarihi, kar_suresi, isik_tarihi, not_txt
                        ])
                if clean_prod:
                    sheets[ws.title] = [HEADERS] + clean_prod
        if sheets:
            return sheets
    except Exception as e:
        st.sidebar.warning(f"Google Sheets yüklenemedi: {e}")

    return get_clean_empty_sheet_data()

def load_seeds_from_sheets():
    try:
        gc = get_gspread_client()
        sh = gc.open("Mikrofiliz_Verileri")
        try:
            ws = sh.worksheet("Tohum_Stoklari")
            rows = ws.get_all_values()
            if len(rows) > 1:
                seeds = []
                for r in rows[1:]:
                    if len(r) >= 5:
                        seeds.append({
                            "product": str(r[0]).strip(),
                            "brand": str(r[1]).strip(),
                            "stock_gram": safe_float(r[2]),
                            "per_unit_gram": safe_float(r[3]),
                            "manual_used_gram": safe_float(r[4]),
                            "is_active": str(r[5]).strip().lower() in ["true", "1", "evet"] if len(r)>5 else True
                        })
                if seeds:
                    return seeds
        except:
            pass
    except Exception as e:
        pass
        
    default_seeds = []
    for name, _, _, gram_unit in ORIGINAL_PRODUCT_TEMPLATES:
        default_seeds.append({
            "product": name,
            "brand": "Varsayılan Marka",
            "stock_gram": 1000.0,
            "per_unit_gram": float(gram_unit),
            "manual_used_gram": 0.0,
            "is_active": True
        })
    return default_seeds

def save_seeds_to_sheets(seeds_list):
    try:
        gc = get_gspread_client()
        sh = gc.open("Mikrofiliz_Verileri")
        try:
            ws = sh.worksheet("Tohum_Stoklari")
            ws.clear()
        except:
            ws = sh.add_worksheet(title="Tohum_Stoklari", rows="100", cols="10")
            
        seed_headers = ["Ürün Adı", "Marka", "Depo Stoğu (g)", "1 Viyol (g)", "Manuel Harcanan (g)", "Aktif mi"]
        full_rows = [seed_headers]
        for s in seeds_list:
            full_rows.append([
                s["product"], s["brand"], str(s["stock_gram"]), 
                str(s["per_unit_gram"]), str(s["manual_used_gram"]), "True" if s["is_active"] else "False"
            ])
        ws.update("A1", full_rows)
    except Exception as e:
        st.error(f"Tohum Kayıt Hatası: {e}")

def save_data_to_sheets(sheets_data):
    try:
        gc = get_gspread_client()
        sh = gc.open("Mikrofiliz_Verileri")
        
        existing_worksheets = {ws.title: ws for ws in sh.worksheets()}
        
        for s_name, rows in sheets_data.items():
            if s_name == "Tohum_Stoklari": continue
            if s_name in existing_worksheets:
                ws = existing_worksheets[s_name]
                ws.clear()
            else:
                ws = sh.add_worksheet(title=s_name, rows="100", cols="20")
            
            data_rows = rows[1:]
            tot_u, tot_s, tot_f, tot_k = 0, 0, 0, 0
            for r in data_rows:
                tot_u += safe_int(r[1]) if len(r)>1 else 0
                tot_s += safe_int(r[2]) if len(r)>2 else 0
                tot_f += safe_int(r[3]) if len(r)>3 else 0
                tot_k += safe_int(r[4]) if len(r)>4 else 0
            
            tot_row = ["Toplam Üretim", str(tot_u), str(tot_s), str(tot_f), str(tot_k)] + [""] * (len(HEADERS) - 5)
            full_rows = [HEADERS] + data_rows + [tot_row]
            ws.update("A1", full_rows)
            
        for old_title, ws in existing_worksheets.items():
            if old_title != "Tohum_Stoklari" and old_title not in sheets_data and len(sh.worksheets()) > 1:
                sh.del_worksheet(ws)

        st.toast("✅ Veriler Google Sheets'e kaydedildi!", icon="💾")
    except Exception as e:
        st.error(f"Google Sheets Kayıt Hatası: {e}")

def generate_pdf_report(title_text, table_headers, table_rows_data):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=20,
        leftMargin=20,
        topMargin=20,
        bottomMargin=20
    )
    story = []
    
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'ReportTitle',
        fontName='Helvetica-Bold',
        fontSize=16,
        leading=20,
        textColor=colors.HexColor("#1b5e20"),
        alignment=1,
        spaceAfter=15
    )
    
    header_cell_style = ParagraphStyle(
        'HeaderCellStyle',
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=11,
        textColor=colors.HexColor("#1b5e20"),
        alignment=1
    )
    
    body_cell_style = ParagraphStyle(
        'BodyCellStyle',
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        alignment=1
    )
    
    clean_title = fix_tr_chars_for_pdf(f"Mikrofiliz Takip - {title_text}")
    story.append(Paragraph(clean_title, title_style))
    story.append(Spacer(1, 10))
    
    pdf_table_data = [[Paragraph(fix_tr_chars_for_pdf(h), header_cell_style) for h in table_headers]]
    
    for row in table_rows_data:
        pdf_row = [Paragraph(fix_tr_chars_for_pdf(cell), body_cell_style) for cell in row]
        pdf_table_data.append(pdf_row)
        
    t = Table(pdf_table_data, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#e8f5e9')),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('BOTTOMPADDING', (0,0), (-1,0), 6),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#c8e6c9')),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#f1f8e9')])
    ]))
    
    story.append(t)
    doc.build(story)
    buffer.seek(0)
    return buffer

def export_sheets_to_excel(sheets_data):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    
    for s_name, rows in sheets_data.items():
        if s_name == "Tohum_Stoklari": continue
        ws = wb.create_sheet(title=s_name)
        data_rows = rows[1:]
        tot_u, tot_s, tot_f, tot_k = 0, 0, 0, 0
        for r in data_rows:
            tot_u += safe_int(r[1]) if len(r)>1 else 0
            tot_s += safe_int(r[2]) if len(r)>2 else 0
            tot_f += safe_int(r[3]) if len(r)>3 else 0
            tot_k += safe_int(r[4]) if len(r)>4 else 0
        
        tot_row = ["Toplam Üretim", str(tot_u), str(tot_s), str(tot_f), str(tot_k)] + [""] * (len(HEADERS) - 5)
        full_rows = [HEADERS] + data_rows + [tot_row]

        for r_idx, r_vals in enumerate(full_rows, 1):
            for c_idx, val in enumerate(r_vals, 1):
                val_str = str(val).replace("💡", "").strip()
                try:
                    num_v = float(val_str)
                    ws.cell(row=r_idx, column=c_idx, value=int(num_v) if num_v.is_integer() else num_v)
                except:
                    ws.cell(row=r_idx, column=c_idx, value=val_str)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

def reset_single_parti(parti_name):
    if parti_name in st.session_state.sheets:
        rows = st.session_state.sheets[parti_name]
        for r_idx in range(1, len(rows)):
            r = rows[r_idx]
            if len(r) > 0 and r[0] != "Toplam Üretim":
                r[1] = "0"
                r[2] = "0"
                r[3] = "0"
                r[4] = "0"
                r[5] = ""
                r[6] = ""
                r[8] = ""
                r[10] = ""
                r[11] = ""
        save_data_to_sheets(st.session_state.sheets)

if "sheets" not in st.session_state:
    st.session_state.sheets = load_data_from_sheets()

if "archived_partis" not in st.session_state:
    st.session_state.archived_partis = []

if "customers_list" not in st.session_state:
    st.session_state.customers_list = list(DEFAULT_CUSTOMERS)

if "brand_seed_stocks" not in st.session_state:
    st.session_state.brand_seed_stocks = load_seeds_from_sheets()

all_parti_names = [k for k in st.session_state.sheets.keys() if k != "Tohum_Stoklari"]
active_parti_names = [p for p in all_parti_names if p not in st.session_state.archived_partis]

if not active_parti_names:
    first_sheet_rows = list(st.session_state.sheets.values())[0][1:]
    new_rows = [HEADERS]
    for r in first_sheet_rows:
        if len(r)>0 and r[0] != "Toplam Üretim":
            new_rows.append([r[0], "0", "0", "0", "0", "", "", r[7] if len(r)>7 else "10", "", r[9] if len(r)>9 else "3", "", ""])
    st.session_state.sheets["Parti No.1"] = new_rows
    if "Parti No.1" in st.session_state.archived_partis:
        st.session_state.archived_partis.remove("Parti No.1")
    active_parti_names = ["Parti No.1"]

if "active_parti" not in st.session_state or st.session_state.active_parti not in active_parti_names:
    st.session_state.active_parti = active_parti_names[0]

# ⚙️ UYGULAMA İÇİ SİSTEM VE YÖNETİM PANELİ (EXPANDER)
with st.expander("⚙️ Sistem Yönetimi (Ürün / Müşteri / Arşiv / Parti Sıfırlama)", expanded=False):
    tab_prod, tab_cust, tab_arch, tab_sys = st.tabs(["🌱 Ürün Yönetimi", "🚚 Müşteri Yönetimi", "📦 Arşivlenmiş Partiler", "🧹 Parti Sıfırlama"])
    
    with tab_prod:
        st.markdown("#### 🌱 Ürün Listesi ve Süre Düzenleme")
        first_p_rows = st.session_state.sheets[all_parti_names[0]][1:]
        current_prods = [r[0] for r in first_p_rows if len(r)>0 and r[0] != "Toplam Üretim"]
        
        col_p1, col_p2 = st.columns(2)
        with col_p1:
            st.markdown("##### ➕ Yeni Ürün Ekle")
            new_p_name = st.text_input("Ürün Adı", key="add_new_p_name")
            new_p_hasat = st.number_input("Hasat Süresi (gün)", min_value=1, value=10, key="add_new_p_hasat")
            new_p_kar = st.number_input("Karanlıkta Kalma (gün)", min_value=1, value=3, key="add_new_p_kar")
            
            if st.button("🌱 Yeni Ürünü Ekle", use_container_width=True):
                if new_p_name.strip() and new_p_name.strip() not in current_prods:
                    p_name_clean = new_p_name.strip()
                    for p_key in st.session_state.sheets.keys():
                        if p_key == "Tohum_Stoklari": continue
                        st.session_state.sheets[p_key].append([
                            p_name_clean, "0", "0", "0", "0", "", "", str(new_p_hasat), "", str(new_p_kar), "", ""
                        ])
                    st.session_state.brand_seed_stocks.append({
                        "product": p_name_clean,
                        "brand": "Varsayılan Marka",
                        "stock_gram": 1000.0,
                        "per_unit_gram": 20.0,
                        "manual_used_gram": 0.0,
                        "is_active": True
                    })
                    save_data_to_sheets(st.session_state.sheets)
                    save_seeds_to_sheets(st.session_state.brand_seed_stocks)
                    st.success(f"'{p_name_clean}' tüm partilere ve Tohum Sistemine eklendi!")
                    st.rerun()
                else:
                    st.error("Lütfen geçerli veya listede bulunmayan bir ürün adı girin.")
                    
        with col_p2:
            st.markdown("##### ✏️ Ürün Sürelerini Güncelle / Sil")
            edit_p_select = st.selectbox("Düzenlenecek / Silinecek Ürün", current_prods, key="edit_p_select")
            ref_row = next((r for r in first_p_rows if r[0] == edit_p_select), None)
            cur_hasat = safe_int(ref_row[7]) if ref_row and len(ref_row)>7 else 10
            cur_kar = safe_int(ref_row[9]) if ref_row and len(ref_row)>9 else 3
            
            up_hasat = st.number_input("Yeni Hasat Süresi (gün)", min_value=1, value=cur_hasat, key="up_p_hasat")
            up_kar = st.number_input("Yeni Karanlık Süresi (gün)", min_value=1, value=cur_kar, key="up_p_kar")
            
            c_btn1, c_btn2 = st.columns(2)
            with c_btn1:
                if st.button("💾 Süreleri Güncelle", use_container_width=True):
                    for p_key in st.session_state.sheets.keys():
                        if p_key == "Tohum_Stoklari": continue
                        for r in st.session_state.sheets[p_key][1:]:
                            if len(r)>0 and r[0] == edit_p_select:
                                r[7] = str(up_hasat)
                                r[9] = str(up_kar)
                    save_data_to_sheets(st.session_state.sheets)
                    st.success(f"'{edit_p_select}' ürününün süreleri güncellendi!")
                    st.rerun()
            with c_btn2:
                if st.button("🗑️ Ürünü Sil", use_container_width=True):
                    for p_key in st.session_state.sheets.keys():
                        if p_key == "Tohum_Stoklari": continue
                        st.session_state.sheets[p_key] = [r for r in st.session_state.sheets[p_key] if len(r)==0 or r[0] != edit_p_select]
                    st.session_state.brand_seed_stocks = [s for s in st.session_state.brand_seed_stocks if s["product"] != edit_p_select]
                    save_data_to_sheets(st.session_state.sheets)
                    save_seeds_to_sheets(st.session_state.brand_seed_stocks)
                    st.success(f"'{edit_p_select}' sistemden silindi!")
                    st.rerun()

    with tab_cust:
        st.markdown("#### 🚚 Müşteri / Satış Noktası Listesi")
        col_c1, col_c2 = st.columns(2)
        with col_c1:
            st.markdown("##### ➕ Yeni Müşteri Ekle")
            add_c_name = st.text_input("Müşteri / Restoran Adı", key="tab_add_cust_input")
            if st.button("➕ Müşteriyi Listeye Ekle", use_container_width=True):
                if add_c_name.strip() and add_c_name.strip() not in st.session_state.customers_list:
                    st.session_state.customers_list.append(add_c_name.strip())
                    st.success(f"'{add_c_name.strip()}' müşteri listesine eklendi!")
                    st.rerun()
                    
        with col_c2:
            st.markdown("##### 🗑️ Müşteri Sil")
            del_c_select = st.selectbox("Silinecek Müşteri", st.session_state.customers_list, key="tab_del_cust_select")
            if st.button("🗑️ Müşteriyi Sil", use_container_width=True):
                st.session_state.customers_list.remove(del_c_select)
                st.success(f"'{del_c_select}' müşteri listesinden çıkarıldı!")
                st.rerun()

    with tab_arch:
        st.markdown("#### 📦 Arşivdeki Partiler")
        if st.session_state.archived_partis:
            unarchive_target = st.selectbox("Arşivden Çıkarılacak Parti Seçin", st.session_state.archived_partis, key="unarchive_select")
            if st.button("🔄 Partiyi Arşivden Çıkar (Aktif Yap)", use_container_width=True):
                st.session_state.archived_partis.remove(unarchive_target)
                st.session_state.active_parti = unarchive_target
                st.success(f"'{unarchive_target}' başarıyla aktif partilere geri yüklendi!")
                st.rerun()
        else:
            st.info("Henüz arşive kaldırılmış bir parti bulunmuyor.")

    with tab_sys:
        st.markdown("#### 🧹 Parti Bazlı Veri Sıfırlama")
        target_reset_parti = st.selectbox("Sıfırlanacak Partiyi Seçin", all_parti_names, key="tab_reset_parti_select")
        if st.button(f"🚨 '{target_reset_parti}' Verilerini Sıfırla (0 Yap)", use_container_width=True, type="primary"):
            reset_single_parti(target_reset_parti)
            st.success(f"'{target_reset_parti}' partisine ait tüm veriler sıfırlandı!")
            st.rerun()

st.markdown("---")

# 🛒 MARKALI TOHUM STOK VE ÜRETİM KAPASİTESİ TAKİP MODÜLÜ
with st.expander("🛒 Markalı Tohum Stok Miktarları ve Üretim Kapasitesi Takibi", expanded=True):
    total_production_by_product = {}
    for p_name in active_parti_names:
        for r in st.session_state.sheets.get(p_name, [])[1:]:
            if len(r) > 0 and r[0] != "Toplam Üretim":
                u_n = r[0]
                total_production_by_product[u_n] = total_production_by_product.get(u_n, 0) + safe_int(r[1])

    first_p_rows = st.session_state.sheets[all_parti_names[0]][1:]
    all_current_prods = [r[0] for r in first_p_rows if len(r)>0 and r[0] != "Toplam Üretim"]

    col_st1, col_st2 = st.columns([1.2, 2])

    with col_st1:
        st.markdown("#### 🏷️ Markalı Tohum Ekle / Düzenle / Sil")
        target_prod_for_seed = st.selectbox("İlgili Ürün", all_current_prods, key="seed_target_prod_select")
        
        existing_brands_for_prod = [s for s in st.session_state.brand_seed_stocks if s["product"] == target_prod_for_seed]
        brand_names_list = [s["brand"] for s in existing_brands_for_prod] if existing_brands_for_prod else ["Varsayılan Marka"]
        
        selected_existing_brand = st.selectbox("Mevcut Marka/Paket Seç (Düzenleme veya Silme İçin)", ["➕ Yeni Marka Ekle..."] + brand_names_list, key="seed_brand_select_exist")

        if selected_existing_brand == "➕ Yeni Marka Ekle...":
            seed_brand_name = st.text_input("Marka / Paket Adı", value="A Marka", key="seed_brand_name_input_new")
            seed_stock_gram = st.number_input("Depo Stoğu (Gram)", min_value=0.0, value=1000.0, step=100.0, key="seed_stock_gram_input_new")
            seed_unit_gram = st.number_input("1 Viyol İçin Harcanan (Gram)", min_value=0.1, value=20.0, step=1.0, key="seed_unit_gram_input_new")
            seed_manual_used = st.number_input("Ekstra / Manuel Harcanan Tohum (Gram)", min_value=0.0, value=0.0, step=10.0, key="seed_manual_used_input_new")
            set_as_active = st.checkbox("Üretimde Varsayılan Tohum Olarak Kullan", value=True, key="seed_set_active_check_new")

            if st.button("💾 Yeni Markalı Tohumu Kaydet", use_container_width=True):
                if seed_brand_name.strip():
                    if set_as_active:
                        for other in st.session_state.brand_seed_stocks:
                            if other["product"] == target_prod_for_seed:
                                other["is_active"] = False
                    st.session_state.brand_seed_stocks.append({
                        "product": target_prod_for_seed,
                        "brand": seed_brand_name.strip(),
                        "stock_gram": seed_stock_gram,
                        "per_unit_gram": seed_unit_gram,
                        "manual_used_gram": seed_manual_used,
                        "is_active": set_as_active
                    })
                    save_seeds_to_sheets(st.session_state.brand_seed_stocks)
                    st.success(f"'{target_prod_for_seed}' için {seed_brand_name.strip()} tohumu eklendi ve Google Sheets'e kaydedildi!")
                    st.rerun()
        else:
            curr_brand_obj = next((s for s in existing_brands_for_prod if s["brand"] == selected_existing_brand), None)
            cur_stk = curr_brand_obj["stock_gram"] if curr_brand_obj else 1000.0
            cur_unit = curr_brand_obj["per_unit_gram"] if curr_brand_obj else 20.0
            cur_man = curr_brand_obj.get("manual_used_gram", 0.0) if curr_brand_obj else 0.0
            cur_act = curr_brand_obj["is_active"] if curr_brand_obj else True

            seed_stock_gram = st.number_input("Depo Stoğu (Gram)", min_value=0.0, value=cur_stk, step=100.0, key="seed_stock_gram_input_edit")
            seed_unit_gram = st.number_input("1 Viyol İçin Harcanan (Gram)", min_value=0.1, value=cur_unit, step=1.0, key="seed_unit_gram_input_edit")
            seed_manual_used = st.number_input("Ekstra / Manuel Harcanan Tohum (Gram)", min_value=0.0, value=cur_man, step=10.0, key="seed_manual_used_input_edit")
            set_as_active = st.checkbox("Üretimde Varsayılan Tohum Olarak Kullan", value=cur_act, key="seed_set_active_check_edit")

            c_sd1, c_sd2 = st.columns(2)
            with c_sd1:
                if st.button("💾 Güncelle", use_container_width=True):
                    curr_brand_obj["stock_gram"] = seed_stock_gram
                    curr_brand_obj["per_unit_gram"] = seed_unit_gram
                    curr_brand_obj["manual_used_gram"] = seed_manual_used
                    if set_as_active:
                        for other in st.session_state.brand_seed_stocks:
                            if other["product"] == target_prod_for_seed:
                                other["is_active"] = False
                        curr_brand_obj["is_active"] = True
                    else:
                        curr_brand_obj["is_active"] = False
                    save_seeds_to_sheets(st.session_state.brand_seed_stocks)
                    st.success(f"'{selected_existing_brand}' tohumu güncellendi ve Google Sheets'e kaydedildi!")
                    st.rerun()
            with c_sd2:
                if st.button("🗑️ Tohumu Sil", use_container_width=True):
                    st.session_state.brand_seed_stocks = [s for s in st.session_state.brand_seed_stocks if not (s["product"] == target_prod_for_seed and s["brand"] == selected_existing_brand)]
                    save_seeds_to_sheets(st.session_state.brand_seed_stocks)
                    st.success(f"'{selected_existing_brand}' tohum kaydı silindi!")
                    st.rerun()

        st.markdown("---")
        # 🟢 YENİ EKLENEN HARCANAN TOHUMLARI 0 LAMA BUTONU
        if st.button("🔄 Harcanan Tohum Gramajlarını Sıfırla (0 Yap)", use_container_width=True):
            for s in st.session_state.brand_seed_stocks:
                s["manual_used_gram"] = 0.0
            save_seeds_to_sheets(st.session_state.brand_seed_stocks)
            st.success("Tüm tohum stoklarındaki harcanan gramajlar sıfırlandı!")
            st.rerun()

    with col_st2:
        col_hdr_s1, col_hdr_s2 = st.columns([2.5, 1.5])
        with col_hdr_s1:
            st.markdown("#### 📋 Depodaki Tohum Stokları ve Üretim Potansiyeli")
        with col_hdr_s2:
            only_active_seeds = st.checkbox("⭐ Sadece Kullanılan Tohumları Göster", value=False, key="filter_only_active_seeds")

        seed_table_rows = []
        seed_pdf_rows = []

        for item in st.session_state.brand_seed_stocks:
            p_name = item["product"]
            b_name = item["brand"]
            tot_g = item["stock_gram"]
            u_g = item["per_unit_gram"]
            man_g = item.get("manual_used_gram", 0.0)
            is_act = item["is_active"]

            if only_active_seeds and not is_act:
                continue

            used_g = man_g
            if is_act:
                p_cnt = total_production_by_product.get(p_name, 0)
                used_g += (p_cnt * u_g)

            rem_g = max(0.0, tot_g - used_g)
            possible_viyol = int(rem_g // u_g) if u_g > 0 else 0

            brand_label = f"{'⭐ (Kullanılan) ' if is_act else ''}{b_name}"

            seed_table_rows.append({
                "Üretilen Ürün": p_name,
                "Tohum Markası": brand_label,
                "Depo Stoğu": f"{tot_g:.0f} g",
                "1 Viyol (g)": f"{u_g:.1f} g",
                "Harcanan Toplam": f"{used_g:.0f} g",
                "Kalan Tohum": f"{rem_g:.0f} g",
                "🌱 Yapılabilecek Üretim": f"{possible_viyol} Adet Viyol"
            })
            seed_pdf_rows.append([p_name, brand_label, f"{tot_g:.0f}g", f"{u_g:.1f}g", f"{used_g:.0f}g", f"{rem_g:.0f}g", f"{possible_viyol} Viyol"])

        if seed_pdf_rows:
            pdf_seed = generate_pdf_report(
                "Tohum Stok ve Uretim Kapasitesi Raporu",
                ["Urun", "Marka", "Depo Stogu", "1 Viyol (g)", "Harcanan", "Kalan Tohum", "Uretim Potansiyeli"],
                seed_pdf_rows
            )
            st.download_button(
                "🖨️ Tohum Raporu PDF",
                data=pdf_seed,
                file_name="Tohum_Stok_Kapasite_Raporu.pdf",
                mime="application/pdf",
                use_container_width=True
            )

        if seed_table_rows:
            df_seeds = pd.DataFrame(seed_table_rows)
            st.dataframe(df_seeds, hide_index=True, use_container_width=True)

st.markdown("---")

# YAN MENÜ: Aktif Parti Seçimi, Arşive Kaldırma, Excel & Oturum Kapatma
with st.sidebar:
    st.header("📂 Parti Yönetimi")
    
    selected_sidebar_parti = st.selectbox(
        "Aktif Parti Seçin", 
        active_parti_names, 
        index=active_parti_names.index(st.session_state.active_parti) if st.session_state.active_parti in active_parti_names else 0,
        key="sb_parti_select"
    )
    if selected_sidebar_parti != st.session_state.active_parti:
        st.session_state.active_parti = selected_sidebar_parti
        st.rerun()

    col_p1, col_p2 = st.columns(2)
    with col_p1:
        if st.button("➕ Yeni Parti"):
            nums = [int(p.replace("Parti No.", "").strip()) for p in all_parti_names if p.startswith("Parti No.")]
            next_num = max(nums) + 1 if nums else 1
            new_name = f"Parti No.{next_num}"
            
            first_sheet_rows = st.session_state.sheets[all_parti_names[0]][1:]
            new_rows = [HEADERS]
            for r in first_sheet_rows:
                if len(r)>0 and r[0] != "Toplam Üretim":
                    new_rows.append([r[0], "0", "0", "0", "0", "", "", r[7] if len(r)>7 else "10", "", r[9] if len(r)>9 else "3", "", ""])
            st.session_state.sheets[new_name] = new_rows
            st.session_state.active_parti = new_name
            save_data_to_sheets(st.session_state.sheets)
            st.rerun()

    with col_p2:
        if len(active_parti_names) > 1:
            if st.button("🗑️ Partiyi Sil"):
                del st.session_state.sheets[st.session_state.active_parti]
                active_parti_names = [p for p in list(st.session_state.sheets.keys()) if p not in st.session_state.archived_partis and p != "Tohum_Stoklari"]
                st.session_state.active_parti = active_parti_names[0]
                save_data_to_sheets(st.session_state.sheets)
                st.rerun()

    st.markdown("<br>", unsafe_allow_html=True)

    if st.button(f"📦 '{st.session_state.active_parti}' Partisini Arşive Kaldır", use_container_width=True):
        if st.session_state.active_parti not in st.session_state.archived_partis:
            st.session_state.archived_partis.append(st.session_state.active_parti)
            st.success(f"'{st.session_state.active_parti}' arşive kaldırıldı!")
            active_parti_names = [p for p in all_parti_names if p not in st.session_state.archived_partis]
            if active_parti_names:
                st.session_state.active_parti = active_parti_names[0]
            st.rerun()

    st.markdown("---")
    st.header("💾 Tam Sistem Yedeği (Kayıpsız)")
    
    full_backup_data = {
        "sheets": st.session_state.sheets,
        "archived_partis": st.session_state.archived_partis,
        "customers_list": st.session_state.customers_list,
        "brand_seed_stocks": st.session_state.brand_seed_stocks
    }
    json_backup_str = json.dumps(full_backup_data, ensure_ascii=False, indent=2)
    
    st.download_button(
        label="💾 Tam Yedeği İndir (JSON)",
        data=json_backup_str,
        file_name=f"Mikrofiliz_Tam_Sistem_Yedegi_{datetime.date.today().strftime('%d_%m_%Y')}.json",
        mime="application/json",
        use_container_width=True
    )

    uploaded_json_backup = st.file_uploader("📤 Sistem Yedeği Yükle (.json)", type=["json"])
    if uploaded_json_backup is not None:
        if st.button("🔄 Yedeği Sisteme Yükle", use_container_width=True):
            try:
                backup_obj = json.load(uploaded_json_backup)
                if "sheets" in backup_obj:
                    st.session_state.sheets = backup_obj["sheets"]
                    if "archived_partis" in backup_obj: st.session_state.archived_partis = backup_obj["archived_partis"]
                    if "customers_list" in backup_obj: st.session_state.customers_list = backup_obj["customers_list"]
                    if "brand_seed_stocks" in backup_obj: st.session_state.brand_seed_stocks = backup_obj["brand_seed_stocks"]
                    
                    save_data_to_sheets(st.session_state.sheets)
                    save_seeds_to_sheets(st.session_state.brand_seed_stocks)
                    st.success("Tüm sistem yedeği başarıyla geri yüklendi ve kaydedildi!")
                    st.rerun()
            except Exception as e:
                st.error(f"Yedek yükleme hatası: {e}")

    st.markdown("---")
    st.header("🔒 Oturum Yönetimi")
    if st.button("🔒 Oturumu Kapat", use_container_width=True):
        st.session_state.authenticated = False
        st.rerun()

    st.markdown("---")
    st.header("📥 Excel Olarak İndir")
    excel_buffer = export_sheets_to_excel(st.session_state.sheets)
    st.download_button(
        label="📊 Güncel Verileri Excel Olarak İndir",
        data=excel_buffer,
        file_name=f"Mikrofiliz_Uretim_Takip_{datetime.date.today().strftime('%d_%m_%Y')}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True
    )

# 🔔 DÜN / BUGÜN / YARIN BİLDİRİM VE GÖREV PANELİ
today = datetime.date.today()
yesterday = today - datetime.timedelta(days=1)
tomorrow = today + datetime.timedelta(days=1)

str_yesterday = yesterday.strftime("%d.%m.%Y")
str_today = today.strftime("%d.%m.%Y")
str_tomorrow = tomorrow.strftime("%d.%m.%Y")

dun_islem = []
bugun_islem = []
yarin_islem = []

for p_name in active_parti_names:
    p_rows = st.session_state.sheets.get(p_name, [])
    for r in p_rows[1:]:
        if len(r) > 0:
            u_name = r[0]
            u_adet = safe_int(r[1])
            ekim_str = r[6]
            hasat_str = r[8]
            isik_str = r[10]

            if u_adet > 0 and ekim_str:
                isik_dt = parse_tr_date(isik_str)
                hasat_dt = parse_tr_date(hasat_str)

                if hasat_dt == yesterday:
                    dun_islem.append(f"✂️ **{u_name}** ({p_name}) - {u_adet} Adet Hasat Edilmeliydi")
                elif isik_dt == yesterday:
                    dun_islem.append(f"💡 **{u_name}** ({p_name}) - {u_adet} Adet Işığa Çıkarılmalıydı")

                if hasat_dt == today:
                    bugun_islem.append(f"✂️ **{u_name}** ({p_name}) - {u_adet} Adet Hasat Edilecek")
                elif isik_dt == today:
                    bugun_islem.append(f"💡 **{u_name}** ({p_name}) - {u_adet} Adet Işığa Çıkarılacak")

                if hasat_dt == tomorrow:
                    yarin_islem.append(f"✂️ **{u_name}** ({p_name}) - {u_adet} Adet Hasat Edilecek")
                elif isik_dt == tomorrow:
                    yarin_islem.append(f"💡 **{u_name}** ({p_name}) - {u_adet} Adet Işığa Çıkarılacak")

if dun_islem or bugun_islem or yarin_islem:
    st.markdown("### 🔔 Görev & Takip Paneli")
    col_d, col_b, col_y = st.columns(3)
    
    with col_d:
        if dun_islem:
            st.error(f"🔴 **DÜN ({str_yesterday}):**\n\n" + "\n\n".join(dun_islem))
        else:
            st.info(f"🔴 **DÜN ({str_yesterday}):**\nGeciken işlem yok.")

    with col_b:
        if bugun_islem:
            st.warning(f"🟡 **BUGÜN ({str_today}):**\n\n" + "\n\n".join(bugun_islem))
        else:
            st.success(f"🟡 **BUGÜN ({str_today}):**\nAcil işlem yok.")

    with col_y:
        if yarin_islem:
            st.info(f"🔵 **YARIN ({str_tomorrow}):**\n\n" + "\n\n".join(yarin_islem))
        else:
            st.info(f"🔵 **YARIN ({str_tomorrow}):**\nPlanlı işlem yok.")
            
    st.markdown("---")

# 🛍️ SATIŞ YAP VE MÜŞTERİ KAYIT MODÜLÜ
with st.expander("🛍️ Satış Yap ve Müşteri Yönetimi", expanded=True):
    col_hdr_sat1, col_hdr_sat2 = st.columns([3, 1])
    with col_hdr_sat1:
        st.markdown("#### 🛍️ Satış İşlemleri ve Kayıt")
        
    col_s1, col_s2, col_s3 = st.columns([1, 1, 1.2])
    
    with col_s1:
        s_parti = st.selectbox("📂 Satış Yapılacak Parti", active_parti_names, key="sales_parti_select")
        
    s_rows = st.session_state.sheets.get(s_parti, [])[1:]
    s_prod_names = [r[0] for r in s_rows if len(r) > 0 and r[0] != "Toplam Üretim"]
    
    with col_s2:
        s_urun = st.selectbox("🌱 Satılacak Ürün", s_prod_names, key="sales_urun_select")
        
    s_idx = next((i for i, r in enumerate(s_rows) if r[0] == s_urun), None)
    
    if s_idx is not None:
        curr_s_row = s_rows[s_idx]
        u_miktari = safe_int(curr_s_row[1])
        mevcut_satilan = safe_int(curr_s_row[2])
        mevcut_fire = safe_int(curr_s_row[3])
        mevcut_kalan = max(0, u_miktari - mevcut_satilan - mevcut_fire)
        ekim_tarihi = curr_s_row[6] if curr_s_row[6] else "-"
        hasat_tarihi = curr_s_row[8] if curr_s_row[8] else "-"
        
        with col_s3:
            st.info(f"📌 **{s_parti} - {s_urun}**\n\nEkim: **{ekim_tarihi}** | Hasat: **{hasat_tarihi}**\n\n**Mevcut Stok:** {mevcut_kalan} Adet (Satılan: {mevcut_satilan} | Fire: {mevcut_fire})")
            
        c_sale1, c_sale2, c_sale3 = st.columns([1, 1.5, 1])
        
        with c_sale1:
            satilacak_adet = st.number_input("Satış Yapılan Adet", min_value=1, max_value=max(1, mevcut_kalan) if mevcut_kalan > 0 else 1, value=1, key="satilacak_adet_input")
            
        customer_options = st.session_state.customers_list
        
        with c_sale2:
            final_customer = st.selectbox("Müşteri / Satış Noktası", customer_options, key="customer_select_option")
            
        with c_sale3:
            st.markdown("<br>", unsafe_allow_html=True)
            if st.button("🛍️ Satışı Onayla ve Tabloya İşle", use_container_width=True):
                if mevcut_kalan < satilacak_adet:
                    st.error("Yetersiz stok! Satış adedi kalan stoktan fazla olamaz.")
                elif not final_customer:
                    st.error("Lütfen geçerli bir müşteri seçin.")
                else:
                    yeni_satilan_toplam = mevcut_satilan + satilacak_adet
                    yeni_kalan = max(0, u_miktari - yeni_satilan_toplam - mevcut_fire)
                    
                    eski_musteri_metni = curr_s_row[5]
                    yeni_kayit_metni = f"{final_customer} ({satilacak_adet} Adet)"
                    if eski_musteri_metni.strip():
                        yeni_musteri_metni = f"{eski_musteri_metni}, {yeni_kayit_metni}"
                    else:
                        yeni_musteri_metni = yeni_kayit_metni
                    
                    curr_s_row[2] = str(yeni_satilan_toplam)
                    curr_s_row[4] = str(yeni_kalan)
                    curr_s_row[5] = yeni_musteri_metni
                    
                    st.session_state.sheets[s_parti][s_idx + 1] = curr_s_row
                    save_data_to_sheets(st.session_state.sheets)
                    st.success(f"✅ Satış Başarılı! {final_customer} müşterisine {satilacak_adet} adet {s_urun} satıldı.")
                    st.rerun()

        sales_pdf_rows = []
        for p_k in active_parti_names:
            for r in st.session_state.sheets[p_k][1:]:
                if len(r)>5 and r[5].strip() and safe_int(r[2])>0:
                    sales_pdf_rows.append([p_k, r[0], r[2], r[5], r[6], r[8]])

        with col_hdr_sat2:
            if sales_pdf_rows:
                pdf_sales = generate_pdf_report(
                    "Satis ve Musteri Dagilim Raporu",
                    ["Parti", "Urun", "Satilan Adet", "Musteri / Satis Noktasi", "Ekim Tarihi", "Hasat Tarihi"],
                    sales_pdf_rows
                )
                st.download_button(
                    "🖨️ Satış Raporu PDF",
                    data=pdf_sales,
                    file_name="Satis_ve_Musteri_Raporu.pdf",
                    mime="application/pdf",
                    use_container_width=True
                )

st.markdown("---")

# 📊 DETAYLI VERİMLİLİK, SATIŞ VE FİRE (ÇÖP) RAPORLAMA MODÜLÜ (TARİH FİLTRELİ)
with st.expander("📊 Üretim, Satış, Çöp (Fire) ve Müşteri Analiz Raporu (Aktif Partiler)", expanded=False):
    col_rep_hdr1, col_rep_hdr2 = st.columns([3, 1])
    
    with col_rep_hdr1:
        st.markdown("#### 📅 Tarih Aralığı Filtresi (Ekim Tarihine Göre)")
        
    col_dr1, col_dr2, col_dr3 = st.columns([1, 1, 1])
    
    with col_dr1:
        start_date_filter = st.date_input("Başlangıç Tarihi", datetime.date.today() - datetime.timedelta(days=30), format="DD.MM.YYYY", key="report_start_date")
    with col_dr2:
        end_date_filter = st.date_input("Bitiş Tarihi", datetime.date.today() + datetime.timedelta(days=30), format="DD.MM.YYYY", key="report_end_date")
    with col_dr3:
        use_date_filter = st.checkbox("Tarih Filtresini Uygula", value=False, key="report_use_date_filter")

    report_list = []
    customer_sales = {}
    tot_global_u, tot_global_s, tot_global_f, tot_global_k = 0, 0, 0, 0
    analysis_pdf_rows = []

    for p_name in active_parti_names:
        p_rows = st.session_state.sheets.get(p_name, [])
        for r in p_rows[1:]:
            if len(r) > 0 and r[0] != "Toplam Üretim":
                u_name = r[0]
                u_adet = safe_int(r[1])
                s_cikan = safe_int(r[2])
                fire = safe_int(r[3])
                kalan = safe_int(r[4])
                musteri = r[5]
                ekim_dt = parse_tr_date(r[6])

                if use_date_filter:
                    if not ekim_dt or not (start_date_filter <= ekim_dt <= end_date_filter):
                        continue

                tot_global_u += u_adet
                tot_global_s += s_cikan
                tot_global_f += fire
                tot_global_k += kalan

                if u_adet > 0 or s_cikan > 0 or fire > 0 or kalan > 0:
                    report_list.append({
                        "Parti": p_name,
                        "Ürün": u_name,
                        "Üretim": u_adet,
                        "Satılan": s_cikan,
                        "Çöpe Atılan (Fire)": fire,
                        "Kalan": kalan
                    })
                    analysis_pdf_rows.append([p_name, u_name, str(u_adet), str(s_cikan), str(fire), str(kalan), r[6], r[8]])

                if musteri and s_cikan > 0:
                    for m_item in musteri.split(","):
                        m_item = m_item.strip()
                        if "(" in m_item and ")" in m_item:
                            m_name = m_item.split("(")[0].strip()
                            m_val = safe_int(m_item.split("(")[1].replace("Adet", "").replace(")", "").strip())
                            if m_name and not m_name.isdigit():
                                customer_sales[m_name] = customer_sales.get(m_name, 0) + m_val
                        elif m_item and not m_item.isdigit():
                            customer_sales[m_item] = customer_sales.get(m_item, 0) + s_cikan

    with col_rep_hdr2:
        if analysis_pdf_rows:
            pdf_analysis = generate_pdf_report(
                "Uretim, Satis ve Fire Analiz Raporu",
                ["Parti", "Urun", "Uretim", "Satilan", "Fire", "Kalan", "Ekim", "Hasat"],
                analysis_pdf_rows
            )
            st.download_button(
                "🖨️ Analiz Raporu PDF",
                data=pdf_analysis,
                file_name="Uretim_Satis_Fire_Analiz_Raporu.pdf",
                mime="application/pdf",
                use_container_width=True
            )

    m1, m2, m3, m4 = st.columns(4)
    m1.metric("🌱 Seçilen Aralıkta Üretim", f"{tot_global_u} Adet")
    m2.metric("🛍️ Satılan / Çıkan", f"{tot_global_s} Adet")
    m3.metric("🗑️ Çöpe Atılan (Fire)", f"{tot_global_f} Adet")
    m4.metric("🛒 Depodaki Kalan", f"{tot_global_k} Adet")

    col_g1, col_g2 = st.columns([2, 1])

    with col_g1:
        if report_list:
            df_report = pd.DataFrame(report_list)
            df_grouped = df_report.groupby("Ürün")[["Üretim", "Satılan", "Çöpe Atılan (Fire)", "Kalan"]].sum().reset_index()
            st.markdown("#### 📈 Ürün Bazlı Üretim, Satış, Fire ve Stok Dağılımı")
            st.bar_chart(df_grouped.set_index("Ürün"))
        else:
            st.info("Seçilen tarih aralığında gösterilecek veri bulunamadı.")

    with col_g2:
        st.markdown("#### 🚚 Müşteri / Satış Noktası Dağılımı")
        if customer_sales:
            df_cust = pd.DataFrame(list(customer_sales.items()), columns=["Müşteri / Yer", "Satılan Adet"]).sort_values(by="Satılan Adet", ascending=False)
            st.dataframe(df_cust, hide_index=True, use_container_width=True)
        else:
            st.info("Seçilen aralıkta müşteri satış kaydı girilmedi.")

st.markdown("---")

# ÜRÜN BAZLI VE EKİLENLER FİLTRELEME MODU (PDF DESTEKLİ)
selected_parti = st.session_state.active_parti
rows = st.session_state.sheets.get(selected_parti, [])

if rows:
    prod_rows = rows[1:]
    all_product_names = sorted(list(set([r[0] for p_k in active_parti_names for r in st.session_state.sheets[p_k][1:] if len(r)>0])))

    col_f1, col_f2 = st.columns([2, 1])
    with col_f1:
        selected_filter_product = st.selectbox(
            "🔍 Ürün Bazlı İnceleme (Aktif Partiler Genelinde):",
            ["Tüm Tabloyu Göster (Aktif Parti)"] + all_product_names
        )
    with col_f2:
        st.markdown("<br>", unsafe_allow_html=True)
        only_planted = st.checkbox("🌱 Sadece Ekilen Ürünleri Göster", value=False, key="filter_only_planted")

    if only_planted:
        prod_rows = [r for r in prod_rows if safe_int(r[1]) > 0 or (len(r) > 6 and r[6].strip() != "")]

    if selected_filter_product != "Tüm Tabloyu Göster (Aktif Parti)":
        col_prod_hdr1, col_prod_hdr2 = st.columns([3, 1])
        with col_prod_hdr1:
            st.subheader(f"📊 '{selected_filter_product}' İçin Aktif Partilerin Üretim ve Satış Geçmişi")
        
        product_history = []
        product_pdf_rows = []
        tot_u, tot_s, tot_f, tot_k = 0, 0, 0, 0

        for p_name in active_parti_names:
            p_rows = st.session_state.sheets[p_name]
            for r in p_rows[1:]:
                if len(r) > 0 and r[0] == selected_filter_product:
                    u_adet = safe_int(r[1])
                    s_cikan = safe_int(r[2])
                    fire = safe_int(r[3])
                    kalan = safe_int(r[4])
                    musteri = r[5]
                    ekim_str = r[6]
                    hasat_str = r[8]

                    tot_u += u_adet
                    tot_s += s_cikan
                    tot_f += fire
                    tot_k += kalan

                    isik_dt = parse_tr_date(r[10])
                    hasat_dt = parse_tr_date(hasat_str)

                    durum_badge = "⚪ Ekilmedi"
                    if hasat_dt and hasat_dt <= today:
                        durum_badge = "🟢 HASAT ZAMANI!"
                    elif isik_dt and isik_dt <= today:
                        durum_badge = "🟡 IŞIĞA ÇIKMALI 💡"
                    elif ekim_str:
                        durum_badge = "🩶 KARANLIKTA"

                    product_history.append({
                        "Parti Adı": p_name,
                        "Durum": durum_badge,
                        "Üretim Adeti": u_adet,
                        "Satılan": s_cikan,
                        "Çöpe Atılan (Fire)": fire,
                        "Kalan Ürün": kalan,
                        "Satış Yapılan Yer / Müşteri": musteri,
                        "Ekim Günü": ekim_str,
                        "Hasat Tarihi": hasat_str,
                        "Not": r[11]
                    })
                    product_pdf_rows.append([p_name, durum_badge, str(u_adet), str(s_cikan), str(fire), str(kalan), musteri, ekim_str, hasat_str])

        with col_prod_hdr2:
            if product_pdf_rows:
                pdf_prod = generate_pdf_report(
                    f"{selected_filter_product} Urun Takip Raporu",
                    ["Parti", "Durum", "Uretim", "Satilan", "Fire", "Kalan", "Musteri", "Ekim", "Hasat"],
                    product_pdf_rows
                )
                st.download_button(
                    "🖨️ Ürün Raporu PDF",
                    data=pdf_prod,
                    file_name=f"{selected_filter_product}_Raporu.pdf",
                    mime="application/pdf",
                    use_container_width=True
                )

        if product_history:
            product_history.append({
                "Parti Adı": "📊 GENEL TOPLAM",
                "Durum": "-",
                "Üretim Adeti": tot_u,
                "Satılan": tot_s,
                "Çöpe Atılan (Fire)": tot_f,
                "Kalan Ürün": tot_k,
                "Satış Yapılan Yer / Müşteri": "-",
                "Ekim Günü": "-",
                "Hasat Tarihi": "-",
                "Not": "-"
            })
            st.table(product_history)
        else:
            st.info("Bu ürün için henüz kayıtlı bir üretim verisi bulunamadı.")

    else:
        col_main_hdr1, col_main_hdr2 = st.columns([3, 1])
        with col_main_hdr1:
            st.markdown(f"### 📋 {selected_parti} Üretim Tablosu")
        
        table_data = []
        parti_pdf_rows = []
        tot_u, tot_s, tot_f, tot_k = 0, 0, 0, 0

        for r in prod_rows:
            u_name = r[0]
            u_adet = safe_int(r[1])
            s_cikan = safe_int(r[2])
            fire = safe_int(r[3])
            kalan = safe_int(r[4])
            musteri = r[5]
            ekim_str = r[6]
            hasat_suresi = r[7]
            hasat_str = r[8]
            karanlik_suresi = r[9]
            isik_str = r[10]
            not_str = r[11]

            tot_u += u_adet
            tot_s += s_cikan
            tot_f += fire
            tot_k += kalan

            isik_dt = parse_tr_date(isik_str)
            hasat_dt = parse_tr_date(hasat_str)

            durum_badge = "⚪ Ekilmedi"
            if hasat_dt and hasat_dt <= today:
                durum_badge = "🟢 HASAT ZAMANI!"
            elif isik_dt and isik_dt <= today:
                durum_badge = "🟡 IŞIĞA ÇIKMALI 💡"
            elif ekim_str:
                durum_badge = "🩶 KARANLIKTA"

            table_data.append({
                "Durum": durum_badge,
                "Ürün Adı": u_name,
                "Üretim Adeti": u_adet,
                "Satılan": s_cikan,
                "Çöpe Atılan (Fire)": fire,
                "Kalan Ürün": kalan,
                "Satış Yapılan Yer / Müşteri": musteri,
                "Ekim Günü": ekim_str,
                "Karanlıkta (gün)": karanlik_suresi,
                "Işığa Çıkış Tarihi": isik_str,
                "Hasat Süresi (gün)": hasat_suresi,
                "Hasat Tarihi": hasat_str,
                "Not": not_str
            })
            parti_pdf_rows.append([durum_badge, u_name, str(u_adet), str(s_cikan), str(fire), str(kalan), musteri, ekim_str, hasat_str])

        with col_main_hdr2:
            if parti_pdf_rows:
                pdf_parti = generate_pdf_report(
                    f"{selected_parti} Uretim Tablosu",
                    ["Durum", "Urun", "Uretim", "Satilan", "Fire", "Kalan", "Musteri", "Ekim", "Hasat"],
                    parti_pdf_rows
                )
                st.download_button(
                    "🖨️ Parti Tablosunu PDF Yap",
                    data=pdf_parti,
                    file_name=f"{selected_parti}_Tablosu.pdf",
                    mime="application/pdf",
                    use_container_width=True
                )

        table_data.append({
            "Durum": "📊 TOPLAM",
            "Ürün Adı": "Toplam Üretim",
            "Üretim Adeti": tot_u,
            "Satılan": tot_s,
            "Çöpe Atılan (Fire)": tot_f,
            "Kalan Ürün": tot_k,
            "Satış Yapılan Yer / Müşteri": "-",
            "Ekim Günü": "-",
            "Karanlıkta (gün)": "-",
            "Işığa Çıkış Tarihi": "-",
            "Hasat Süresi (gün)": "-",
            "Hasat Tarihi": "-",
            "Not": "-"
        })

        st.table(table_data)

    # HIZLI DÜZENLEME PANELİ (EKİM & FİRE YÖNETİMİ)
    st.markdown("---")
    st.subheader("⚡ Hızlı İşlem Paneli (Ekim & Fire Yönetimi)")

    col_hp1, col_hp2 = st.columns(2)
    
    with col_hp1:
        selected_urun = st.selectbox("🌱 İşlem Yapılacak Ürün", all_product_names, key="quick_urun_select")

    with col_hp2:
        target_parti = st.selectbox(
            "📂 İşlem Yapılacak Parti",
            active_parti_names,
            index=active_parti_names.index(selected_parti) if selected_parti in active_parti_names else 0,
            key="quick_parti_select"
        )

    target_sheet_rows = st.session_state.sheets.get(target_parti, [])[1:]
    idx = next((i for i, r in enumerate(target_sheet_rows) if r[0] == selected_urun), None)

    if idx is not None:
        curr_r = target_sheet_rows[idx]

        c1, c2, c3 = st.columns(3)
        with c1:
            e_ekim = st.date_input(
                "Ekim Tarihi (Gün/Ay/Yıl)", 
                parse_tr_date(curr_r[6]) if curr_r[6] else datetime.date.today(), 
                format="DD.MM.YYYY",
                key=f"input_ekim_{target_parti}_{selected_urun}"
            )
            e_u_adet = st.number_input(
                "Üretim Adeti", 
                min_value=0, 
                value=safe_int(curr_r[1]),
                key=f"input_u_adet_{target_parti}_{selected_urun}"
            )
        with c2:
            e_fire = st.number_input(
                "Çöpe Atılan (Fire)", 
                min_value=0, 
                value=safe_int(curr_r[3]),
                key=f"input_fire_{target_parti}_{selected_urun}"
            )
            e_kar = st.number_input(
                "Karanlıkta (gün)", 
                min_value=1, 
                value=safe_int(curr_r[9]) if curr_r[9] != "" else 3,
                key=f"input_kar_{target_parti}_{selected_urun}"
            )
        with c3:
            e_has = st.number_input(
                "Hasat Süresi (gün)", 
                min_value=1, 
                value=safe_int(curr_r[7]) if curr_r[7] != "" else 10,
                key=f"input_has_{target_parti}_{selected_urun}"
            )
            e_not = st.text_input(
                "Not / Açıklama", 
                value=curr_r[11],
                key=f"input_not_{target_parti}_{selected_urun}"
            )

        if st.button(f"⚡ {selected_urun} ({target_parti}) Verilerini Kaydet", use_container_width=True):
            e_s_cikan = safe_int(curr_r[2])
            e_musteri = curr_r[5]
            kalan = max(0, e_u_adet - e_s_cikan - e_fire)
            isik_dt = e_ekim + datetime.timedelta(days=int(e_kar))
            hasat_dt = e_ekim + datetime.timedelta(days=int(e_has))

            updated = [
                selected_urun,
                str(e_u_adet),
                str(e_s_cikan),
                str(e_fire),
                str(kalan),
                e_musteri,
                e_ekim.strftime("%d.%m.%Y"),
                str(e_has),
                hasat_dt.strftime("%d.%m.%Y"),
                str(e_kar),
                isik_dt.strftime("%d.%m.%Y"),
                e_not
            ]
            st.session_state.sheets[target_parti][idx + 1] = updated
            save_data_to_sheets(st.session_state.sheets)
            st.success(f"'{selected_urun}' ({target_parti}) güncellendi! Çöpe Atılan: {e_fire} | Kalan Ürün: {kalan}")
            st.rerun()
    else:
        st.warning(f"'{selected_urun}' adlı ürün {target_parti} içerisinde henüz tanımlı değil.")
